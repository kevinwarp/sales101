"""Load local files into LangChain Documents based on file extension."""

from __future__ import annotations

import csv
import logging
from pathlib import Path

from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    UnstructuredPowerPointLoader,
)
from langchain_core.documents import Document

logger = logging.getLogger(__name__)


def load_file(path: Path) -> list[Document]:
    """Return a list of LangChain Documents parsed from *path*.

    Supported extensions: .pdf, .docx, .xlsx, .pptx, .csv
    """
    ext = path.suffix.lower()
    metadata = {"source": str(path), "filename": path.name}

    if ext == ".pdf":
        docs = PyPDFLoader(str(path)).load()

    elif ext == ".docx":
        docs = Docx2txtLoader(str(path)).load()

    elif ext == ".pptx":
        docs = UnstructuredPowerPointLoader(str(path)).load()

    elif ext == ".xlsx":
        docs = _load_xlsx(path)

    elif ext == ".csv":
        docs = _load_csv(path)

    else:
        logger.warning("Unsupported extension %s – skipping %s", ext, path.name)
        return []

    # Ensure consistent metadata on every document
    for doc in docs:
        doc.metadata.update(metadata)

    logger.info("Loaded %d document(s) from %s", len(docs), path.name)
    return docs


# ── Private helpers ──────────────────────────────────────────────────────────


def _load_xlsx(path: Path) -> list[Document]:
    """Read an Excel workbook and return one Document per sheet."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    docs: list[Document] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append("\t".join(str(c) if c is not None else "" for c in row))
        text = "\n".join(rows)
        if text.strip():
            docs.append(Document(page_content=text, metadata={"sheet": sheet_name}))
    return docs


def _load_csv(path: Path) -> list[Document]:
    """Read a CSV file and return a single Document."""
    with open(path, newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        rows = ["\t".join(row) for row in reader]
    text = "\n".join(rows)
    return [Document(page_content=text)] if text.strip() else []
