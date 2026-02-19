"""File-type-specific parsers that produce LangChain Documents."""

from __future__ import annotations

from pathlib import Path

from langchain_community.document_loaders import (
    CSVLoader,
    PyPDFLoader,
    UnstructuredPowerPointLoader,
    UnstructuredWordDocumentLoader,
)
from langchain_core.documents import Document

# Map file suffixes to loader classes.
_LOADER_MAP: dict[str, type] = {
    ".pdf": PyPDFLoader,
    ".docx": UnstructuredWordDocumentLoader,
    ".csv": CSVLoader,
    ".xlsx": None,  # handled separately
    ".pptx": UnstructuredPowerPointLoader,
}


def _load_xlsx(path: Path) -> list[Document]:
    """Load an Excel workbook, one Document per sheet."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    docs: list[Document] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append("\t".join(str(c) if c is not None else "" for c in row))
        text = "\n".join(rows)
        if text.strip():
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": str(path), "sheet": sheet_name},
                )
            )
    return docs


def load_document(path: Path, source_name: str | None = None) -> list[Document]:
    """Load a single file into one or more LangChain Documents."""
    suffix = path.suffix.lower()

    if suffix == ".xlsx":
        docs = _load_xlsx(path)
    elif suffix in _LOADER_MAP:
        loader_cls = _LOADER_MAP[suffix]
        if loader_cls is None:
            raise ValueError(f"Unsupported file type: {suffix}")
        loader = loader_cls(str(path))
        docs = loader.load()
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    # Attach a human-readable source name to every document.
    for doc in docs:
        doc.metadata["source_name"] = source_name or path.name

    return docs
